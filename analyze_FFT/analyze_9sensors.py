# -*- coding: utf-8 -*-
"""分析三设备日志，按 ADC 通道映射到传感器编号，对比每个传感器 Empty vs Full 的 RMS。
   通道映射（A0/IO15/IO16 -> 传感器编号）：
     COM3 (主动):  A0→1,  IO15→4, IO16→5
     COM28(被动):  A0→2,  IO15→6, IO16→7
     COM33(被动):  A0→3,  IO15→8, IO16→9
   用法：
     py analyze_9sensors.py <COM3.txt> <COM28.txt> <COM33.txt>
   输出：同目录下 sensor_analysis.xlsx
"""
import sys, os
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

CHANNEL_MAP = {
    'COM3':  {'A0': 1, 'IO15': 4, 'IO16': 5},
    'COM28': {'A0': 2, 'IO15': 6, 'IO16': 7},
    'COM33': {'A0': 3, 'IO15': 8, 'IO16': 9},
}
FM_COLS = ['B0', 'B1', 'B2', 'B3', 'B4', 'B5', 'B6', 'B7']


def parse_log(path):
    port = os.path.splitext(os.path.basename(path))[0]
    rows = []
    cur_ch = None
    with open(path, encoding='utf-8') as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            if line.startswith('#CHANNEL,'):
                ch = line.split(',', 1)[1].split()[0] if ',' in line else None
                cur_ch = ch
                continue
            if line.startswith('#'):
                continue
            cols = line.split(',')
            if len(cols) < 15 or cols[-1] != '#END':
                continue
            if not cols[0].isdigit():
                continue
            try:
                sid = CHANNEL_MAP.get(port, {}).get(cur_ch)
                rows.append({
                    'sensor': sid, 'port': port, 'channel': cur_ch,
                    'state': cols[1], 'phase': int(cols[2]),
                    'peak': float(cols[3]), 'rms': float(cols[4]), 'std': float(cols[5]),
                    'b0': int(cols[6]), 'b1': int(cols[7]), 'b2': int(cols[8]), 'b3': int(cols[9]),
                    'b4': int(cols[10]), 'b5': int(cols[11]), 'b6': int(cols[12]), 'b7': int(cols[13]),
                })
            except Exception:
                pass
    return rows


def build_excel(df, out_path):
    wb = Workbook()
    hf = Font(bold=True, color='FFFFFF', size=11, name='Arial')
    hfill = PatternFill('solid', fgColor='4472C4')
    halign = Alignment(horizontal='center', vertical='center', wrap_text=True)
    dalign = Alignment(horizontal='center', vertical='center')
    border = Border(left=Side(style='thin', color='D9D9D9'), right=Side(style='thin', color='D9D9D9'),
                    top=Side(style='thin', color='D9D9D9'), bottom=Side(style='thin', color='D9D9D9'))

    def style_header(ws, row, max_col):
        for c in range(1, max_col + 1):
            cell = ws.cell(row=row, column=c)
            cell.font = hf; cell.fill = hfill; cell.alignment = halign; cell.border = border

    def style_data(ws, r1, r2, max_col):
        for r in range(r1, r2 + 1):
            for c in range(1, max_col + 1):
                cell = ws.cell(row=r, column=c)
                cell.alignment = dalign; cell.border = border
                if r % 2 == 0:
                    cell.fill = PatternFill('solid', fgColor='F2F7FB')

    sensors = sorted({s for s in df['sensor'].dropna().unique() if s is not None})
    # ===== Sheet1 传感器对比 =====
    ws = wb.active
    ws.title = '传感器对比'
    hdr = ['传感器', '设备', '通道', 'Empty_RMS均值', 'Empty_Std', 'Full_RMS均值', 'Full_Std',
           '差值(Empty-Full)', '相对区分度%', 'n_Empty', 'n_Full']
    for ci, h in enumerate(hdr, 1):
        ws.cell(row=1, column=ci, value=h)
    style_header(ws, 1, len(hdr))
    r = 2
    for sid in sensors:
        sub = df[df['sensor'] == sid]
        e = sub[sub['state'] == 'Empty']
        fl = sub[sub['state'] == 'Full']
        er = e['rms'].mean() if len(e) else float('nan')
        fr = fl['rms'].mean() if len(fl) else float('nan')
        es = e['rms'].std() if len(e) > 1 else 0.0
        fs = fl['rms'].std() if len(fl) > 1 else 0.0
        diff = er - fr
        rel = (diff / er * 100) if er else float('nan')
        port = sub['port'].iloc[0] if len(sub) else ''
        ch = sub['channel'].dropna().iloc[0] if len(sub['channel'].dropna()) else ''
        vals = [sid, port, ch, round(er, 2), round(es, 2), round(fr, 2), round(fs, 2),
                round(diff, 2), (round(rel, 1) if rel == rel else ''), len(e), len(fl)]
        for ci, v in enumerate(vals, 1):
            ws.cell(row=r, column=ci, value=v)
        r += 1
    style_data(ws, 2, r - 1, len(hdr))
    for ci, h in enumerate(hdr, 1):
        ws.column_dimensions[get_column_letter(ci)].width = max(12, len(str(h)) + 3)

    chart = BarChart()
    chart.type = 'col'
    chart.title = '传感器 Empty vs Full RMS 对比'
    chart.y_axis.title = 'RMS'
    chart.x_axis.title = '传感器编号'
    chart.style = 10
    chart.height = 16; chart.width = 26
    n = len(sensors)
    chart.add_data(Reference(ws, min_col=4, max_col=4, min_row=1, max_row=n + 1), titles_from_data=True)
    chart.add_data(Reference(ws, min_col=6, max_col=6, min_row=1, max_row=n + 1), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n + 1))
    chart.gapWidth = 60
    ws.add_chart(chart, f'M2')

    # ===== Sheet2 FM 频段对比 =====
    ws2 = wb.create_sheet('FM频段对比')
    ws2.cell(row=1, column=1, value='传感器')
    ws2.cell(row=1, column=2, value='频段')
    ws2.cell(row=1, column=3, value='Empty均值')
    ws2.cell(row=1, column=4, value='Full均值')
    ws2.cell(row=1, column=5, value='差值(E-F)')
    style_header(ws2, 1, 5)
    rr = 2
    for sid in sensors:
        sub = df[df['sensor'] == sid]
        e = sub[sub['state'] == 'Empty']
        fl = sub[sub['state'] == 'Full']
        for i, fc in enumerate(FM_COLS):
            ev = e[fc.lower()].mean() if len(e) else 0
            fv = fl[fc.lower()].mean() if len(fl) else 0
            ws2.cell(row=rr, column=1, value=sid)
            ws2.cell(row=rr, column=2, value=fc)
            ws2.cell(row=rr, column=3, value=round(ev, 0))
            ws2.cell(row=rr, column=4, value=round(fv, 0))
            ws2.cell(row=rr, column=5, value=round(ev - fv, 0))
            rr += 1
    style_data(ws2, 2, rr - 1, 5)
    for ci in range(1, 6):
        ws2.column_dimensions[get_column_letter(ci)].width = 14

    # ===== Sheet3 原始数据 =====
    ws3 = wb.create_sheet('原始数据')
    cols = ['sensor', 'port', 'channel', 'state', 'phase', 'peak', 'rms', 'std'] + [c.lower() for c in FM_COLS]
    for ci, c in enumerate(cols, 1):
        ws3.cell(row=1, column=ci, value=c)
    style_header(ws3, 1, len(cols))
    for ri, row in df.iterrows():
        for ci, c in enumerate(cols, 1):
            ws3.cell(row=ri + 2, column=ci, value=row[c])
    style_data(ws3, 2, len(df) + 1, len(cols))
    for ci, c in enumerate(cols, 1):
        ws3.column_dimensions[get_column_letter(ci)].width = max(10, len(c) + 2)

    wb.save(out_path)


def analyze(paths):
    all_rows = []
    for p in paths:
        rows = parse_log(p)
        print(f'{os.path.basename(p)}: {len(rows)} 行', flush=True)
        all_rows.extend(rows)
    df = pd.DataFrame(all_rows)
    print(f'总计 {len(df)} 行, 传感器: {sorted(df["sensor"].dropna().unique())}', flush=True)
    out = os.path.join(os.path.dirname(paths[0]), 'sensor_analysis.xlsx')
    build_excel(df, out)
    print(f'已生成: {out}', flush=True)
    return out, df


def main():
    if len(sys.argv) < 4:
        print('Usage: py analyze_9sensors.py <COM3.txt> <COM28.txt> <COM33.txt>')
        sys.exit(1)
    analyze(sys.argv[1:4])


if __name__ == '__main__':
    main()
